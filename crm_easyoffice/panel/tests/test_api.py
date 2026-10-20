"""Dashboard indicator and export tests. All data here is synthetic."""

from __future__ import annotations

import importlib
from datetime import date
from datetime import datetime
from datetime import timedelta
from decimal import Decimal
from http import HTTPStatus
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from django.contrib.auth.models import Group
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from crm_easyoffice.clientes.models import Cliente
from crm_easyoffice.clientes.tests.factories import ClienteFactory
from crm_easyoffice.clientes.tests.factories import CuentaClienteFactory
from crm_easyoffice.contratos.models import Contrato
from crm_easyoffice.contratos.tests.factories import nuevo_contrato
from crm_easyoffice.documentos.tests.factories import documento_generado
from crm_easyoffice.panel.services import EN_LINEA
from crm_easyoffice.panel.services import calcular_indicadores
from crm_easyoffice.tramites.services import COMPLETADO
from crm_easyoffice.tramites.services import FIRMA
from crm_easyoffice.tramites.services import INICIADO
from crm_easyoffice.tramites.services import PENDIENTE
from crm_easyoffice.tramites.tests.factories import tramite_con_historia
from crm_easyoffice.users.tests.factories import UserFactory
from crm_easyoffice.ventas.models import Venta
from crm_easyoffice.ventas.tests.factories import ServicioFactory
from crm_easyoffice.ventas.tests.factories import venta_pagada

seed = importlib.import_module("crm_easyoffice.core.migrations.0002_seed_roles")

SANTIAGO = ZoneInfo("America/Santiago")
SEPTIEMBRE = {"desde": "2026-09-01", "hasta": "2026-09-30"}
EJECUTIVA = "Ejecutiva de Prueba"
VISTAS = ["panel:indicadores", "panel:exportar"]


def client_for(role: str) -> APIClient:
    user = UserFactory.create()
    user.groups.add(Group.objects.get(name=role))
    api_client = APIClient()
    api_client.force_authenticate(user)
    return api_client


def en_santiago(*partes: int) -> datetime:
    return datetime(*partes, tzinfo=SANTIAGO)


def registrar_cliente(registrado: datetime) -> None:
    cliente = ClienteFactory.create()
    Cliente.objects.filter(pk=cliente.pk).update(created_at=registrado)


def contrato_que_vence(vencimiento: date, **datos) -> Contrato:
    contrato = nuevo_contrato(
        fecha_inicio=date(2026, 1, 5),
        renovacion_automatica=False,
        **datos,
    )
    Contrato.objects.filter(pk=contrato.pk).update(fecha_vencimiento=vencimiento)
    return contrato


def filas_de(libro, hoja: str) -> list[tuple]:
    return list(libro[hoja].iter_rows(values_only=True))


@pytest.fixture
def administrador(db) -> APIClient:
    return client_for(seed.ADMINISTRADOR)


@pytest.fixture
def clientes_de_prueba(db) -> None:
    registrar_cliente(en_santiago(2026, 8, 15, 12))
    registrar_cliente(en_santiago(2026, 9, 1, 0, 30))
    registrar_cliente(en_santiago(2026, 9, 30, 23, 30))
    registrar_cliente(en_santiago(2026, 10, 1, 9))


@pytest.fixture
def operacion(db) -> None:
    """September's activity, read at 30 September."""
    ejecutiva = UserFactory.create(name=EJECUTIVA)
    ejecutiva.groups.add(Group.objects.get(name=seed.EJECUTIVO))

    # Contracts: one in force, one due within 30 days, one expired, a draft.
    contrato_que_vence(date(2027, 1, 4))
    contrato_que_vence(date(2026, 10, 30))
    contrato_que_vence(date(2026, 9, 29))
    nuevo_contrato(fecha_inicio=date(2026, 9, 1), borrador=True)

    # Sales: one through the executive's case, one online, one refunded and
    # one paid after the period.
    domicilio = ServicioFactory.create(
        nombre="Servicio A de prueba",
        precio_base=Decimal(50000),
    )
    firma = ServicioFactory.create(
        nombre="Servicio B de prueba",
        precio_base=Decimal(5000),
    )
    con_ejecutiva = venta_pagada(en_santiago(2026, 9, 1, 0, 30), (domicilio, 1))
    venta_pagada(en_santiago(2026, 9, 30, 23, 30), (domicilio, 1), (firma, 1))
    venta_pagada(
        en_santiago(2026, 9, 15, 12),
        (domicilio, 1),
        estado=Venta.Estado.REEMBOLSADA,
    )
    venta_pagada(en_santiago(2026, 10, 1, 0, 30), (domicilio, 1))

    # Cases: the executive's completed (its contract starts in October), and
    # another waits for its signature.
    tramite_con_historia(
        (en_santiago(2026, 8, 30, 10), INICIADO, None),
        (en_santiago(2026, 9, 1, 10), COMPLETADO, None),
        venta=con_ejecutiva,
        contrato=nuevo_contrato(ejecutivo=ejecutiva),
    )
    en_firma = tramite_con_historia(
        (en_santiago(2026, 9, 20, 10), INICIADO, None),
        (en_santiago(2026, 9, 21, 10), PENDIENTE, FIRMA),
    )
    documento_generado(en_firma, 1, en_santiago(2026, 9, 21, 10))


class TestIndicadores:
    def test_client_counts_are_inclusive_in_local_time(
        self,
        administrador,
        clientes_de_prueba,
    ):
        response = administrador.get(reverse("panel:indicadores"), SEPTIEMBRE)

        assert response.status_code == HTTPStatus.OK
        assert response.data["periodo"] == SEPTIEMBRE
        assert response.data["clientes"] == {"total": 3, "nuevos": 2}

    def test_payload_keeps_the_documented_shape(self, administrador):
        data = administrador.get(reverse("panel:indicadores"), SEPTIEMBRE).data

        assert set(data) == {
            "periodo",
            "fecha_referencia",
            "clientes",
            "servicios",
            "ventas",
            "tramites_pendientes",
            "tramites_pendientes_por_motivo",
            "documentos_pendientes_firma",
            "datos_de_ejemplo",
        }
        assert set(data["servicios"]) == {
            "activos",
            "por_vencer",
            "vencidos",
            "dias_aviso",
        }
        assert set(data["ventas"]) == {
            "moneda",
            "total",
            "por_servicio",
            "por_ejecutivo",
        }
        assert data["datos_de_ejemplo"] == []

    def test_every_figure_is_real(self, administrador, operacion):
        data = administrador.get(reverse("panel:indicadores"), SEPTIEMBRE).data

        assert data["fecha_referencia"] == "2026-09-30"
        assert data["servicios"] == {
            "activos": 2,
            "por_vencer": 1,
            "vencidos": 1,
            "dias_aviso": 30,
        }
        assert data["ventas"] == {
            "moneda": "CLP",
            "total": 105000,
            "por_servicio": [
                {"servicio": "Servicio A de prueba", "monto": 100000, "cantidad": 2},
                {"servicio": "Servicio B de prueba", "monto": 5000, "cantidad": 1},
            ],
            "por_ejecutivo": [
                {"ejecutivo": EJECUTIVA, "monto": 50000, "cantidad": 1},
                {"ejecutivo": EN_LINEA, "monto": 55000, "cantidad": 1},
            ],
        }
        assert data["tramites_pendientes"] == 1
        assert [fila["motivo"] for fila in data["tramites_pendientes_por_motivo"]] == [
            FIRMA,
        ]
        assert data["documentos_pendientes_firma"] == 1

    def test_reference_date_is_today_when_the_period_ends_later(self, administrador):
        hoy = timezone.localdate()
        periodo = {
            "desde": hoy.replace(day=1).isoformat(),
            "hasta": (hoy + timedelta(days=10)).isoformat(),
        }

        data = administrador.get(reverse("panel:indicadores"), periodo).data

        assert data["fecha_referencia"] == hoy.isoformat()

    def test_point_in_time_figures_follow_the_reference_date(self, operacion):
        mitad = calcular_indicadores(
            date(2026, 9, 1),
            date(2026, 9, 30),
            hoy=date(2026, 9, 15),
        )

        assert mitad["fecha_referencia"] == "2026-09-15"
        # The contract expiring on 29 September was due then, not expired.
        assert mitad["servicios"] == {
            "activos": 3,
            "por_vencer": 1,
            "vencidos": 0,
            "dias_aviso": 30,
        }
        assert mitad["tramites_pendientes"] == 0
        assert mitad["documentos_pendientes_firma"] == 0

    def test_queries_do_not_grow_with_the_rows(
        self,
        operacion,
        django_assert_max_num_queries,
    ):
        # Two per client count and per name lookup, one per other figure.
        with django_assert_max_num_queries(10):
            calcular_indicadores(date(2026, 9, 1), date(2026, 9, 30))

    @pytest.mark.parametrize(
        "params",
        [
            {},
            {"desde": "2026-09-01"},
            {"desde": "no-es-fecha", "hasta": "2026-09-30"},
            {"desde": "2026-09-30", "hasta": "2026-09-01"},
        ],
    )
    def test_rejects_an_invalid_period(self, administrador, params):
        response = administrador.get(reverse("panel:indicadores"), params)

        assert response.status_code == HTTPStatus.BAD_REQUEST


class TestAcceso:
    @pytest.mark.parametrize("rol", [seed.ADMINISTRADOR, seed.EJECUTIVO])
    @pytest.mark.parametrize("vista", VISTAS)
    def test_both_staff_roles_see_the_dashboard(self, db, rol, vista):
        response = client_for(rol).get(reverse(vista), SEPTIEMBRE)

        assert response.status_code == HTTPStatus.OK

    @pytest.mark.parametrize("vista", VISTAS)
    def test_portal_clients_and_anonymous_users_are_denied(self, db, vista):
        cliente = APIClient()
        cliente.force_authenticate(CuentaClienteFactory.create().user)

        for api_client in (cliente, APIClient()):
            response = api_client.get(reverse(vista), SEPTIEMBRE)
            assert response.status_code == HTTPStatus.FORBIDDEN


class TestExportar:
    def test_downloads_a_workbook_with_the_indicators(
        self,
        administrador,
        clientes_de_prueba,
        operacion,
    ):
        response = administrador.get(reverse("panel:exportar"), SEPTIEMBRE)

        assert response.status_code == HTTPStatus.OK
        assert response["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml",
        )
        assert response["Content-Disposition"] == (
            'attachment; filename="panel-operativo_2026-09-01_2026-09-30.xlsx"'
        )

        libro = load_workbook(BytesIO(response.content))
        assert libro.sheetnames == [
            "Resumen",
            "Ventas por servicio",
            "Ventas por ejecutivo",
        ]
        filas = {fila[0]: fila[1:] for fila in filas_de(libro, "Resumen") if fila[0]}
        assert filas["Total de clientes"] == (3, None)
        assert filas["Clientes nuevos"] == (2, "En el período")
        assert filas["Servicios activos"] == (2, "Al 30-09-2026")
        assert filas["Servicios por vencer"] == (1, "Próximos 30 días, al 30-09-2026")
        assert filas["Ventas totales (CLP)"] == (
            105000,
            "Pagadas en línea, en el período",
        )
        assert filas["Documentos pendientes de firma"] == (1, "Al 30-09-2026")

    def test_nothing_is_labelled_as_example(self, administrador, operacion):
        response = administrador.get(reverse("panel:exportar"), SEPTIEMBRE)

        libro = load_workbook(BytesIO(response.content))
        celdas = [
            str(valor).lower()
            for hoja in libro.sheetnames
            for fila in filas_de(libro, hoja)
            for valor in fila
            if valor is not None
        ]
        assert not [celda for celda in celdas if "ejemplo" in celda]

    def test_sales_sheets_add_up_to_the_total(self, administrador, operacion):
        indicadores = administrador.get(reverse("panel:indicadores"), SEPTIEMBRE).data
        response = administrador.get(reverse("panel:exportar"), SEPTIEMBRE)

        libro = load_workbook(BytesIO(response.content))
        assert EN_LINEA in [fila[0] for fila in filas_de(libro, "Ventas por ejecutivo")]
        for hoja in ("Ventas por servicio", "Ventas por ejecutivo"):
            total = next(fila for fila in filas_de(libro, hoja) if fila[0] == "Total")
            assert total[2] == indicadores["ventas"]["total"]
