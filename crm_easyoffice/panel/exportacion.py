"""Excel export of the operational dashboard."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import TYPE_CHECKING
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

FORMATO_CLP = '"$"#,##0'
FORMATO_FECHA = "dd-mm-yyyy"
FORMATO_FECHA_HORA = "dd-mm-yyyy hh:mm"
VENTAS_EN_LINEA = "Pagadas en línea, en el período"
NEGRITA = Font(bold=True)


def nombre_archivo(indicadores: dict[str, Any]) -> str:
    periodo = indicadores["periodo"]
    return f"panel-operativo_{periodo['desde']}_{periodo['hasta']}.xlsx"


def libro_indicadores(indicadores: dict[str, Any]) -> bytes:
    """Build the workbook for the payload of ``calcular_indicadores``.

    Sheets: Resumen with the eight RF-14 indicators, then sales by service and
    by executive. Point-in-time rows say the day they were read at.
    """
    libro = Workbook()
    resumen = libro.active
    assert resumen is not None
    resumen.title = "Resumen"
    _escribir_resumen(resumen, indicadores)

    ventas = indicadores["ventas"]
    for titulo, clave, filas in (
        ("Servicio", "servicio", ventas["por_servicio"]),
        ("Ejecutivo", "ejecutivo", ventas["por_ejecutivo"]),
    ):
        hoja = libro.create_sheet(f"Ventas por {clave}")
        _escribir_ventas(hoja, titulo, clave, filas)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _escribir_resumen(hoja: Worksheet, indicadores: dict[str, Any]) -> None:
    periodo = indicadores["periodo"]
    hoja.append(["Panel operativo"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(["Desde", date.fromisoformat(periodo["desde"])])
    hoja.append(["Hasta", date.fromisoformat(periodo["hasta"])])
    hoja.append(["Generado", timezone.localtime().replace(tzinfo=None)])
    for fila in (2, 3):
        hoja.cell(fila, 2).number_format = FORMATO_FECHA
    hoja.cell(4, 2).number_format = FORMATO_FECHA_HORA

    hoja.append([])

    hoja.append(["Indicador", "Valor", "Detalle"])
    for celda in hoja[hoja.max_row]:
        celda.font = NEGRITA

    for etiqueta, valor, clave, detalle in _filas_resumen(indicadores):
        hoja.append([etiqueta, valor, detalle])
        if clave == "ventas":
            hoja.cell(hoja.max_row, 2).number_format = FORMATO_CLP

    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 18
    hoja.column_dimensions["C"].width = 32


def _filas_resumen(indicadores: dict[str, Any]) -> list[tuple[str, int, str, str]]:
    clientes = indicadores["clientes"]
    servicios = indicadores["servicios"]
    al = f"Al {date.fromisoformat(indicadores['fecha_referencia']):%d-%m-%Y}"
    return [
        ("Total de clientes", clientes["total"], "clientes", ""),
        ("Clientes nuevos", clientes["nuevos"], "clientes", "En el período"),
        ("Servicios activos", servicios["activos"], "servicios", al),
        (
            "Servicios por vencer",
            servicios["por_vencer"],
            "servicios",
            f"Próximos {servicios['dias_aviso']} días, {al.lower()}",
        ),
        ("Servicios vencidos", servicios["vencidos"], "servicios", al),
        (
            "Ventas totales (CLP)",
            indicadores["ventas"]["total"],
            "ventas",
            VENTAS_EN_LINEA,
        ),
        (
            "Trámites pendientes",
            indicadores["tramites_pendientes"],
            "tramites_pendientes",
            al,
        ),
        (
            "Documentos pendientes de firma",
            indicadores["documentos_pendientes_firma"],
            "documentos_pendientes_firma",
            al,
        ),
    ]


def _escribir_ventas(
    hoja: Worksheet,
    titulo: str,
    clave: str,
    filas: list[dict[str, Any]],
) -> None:
    hoja.append([titulo, "Cantidad", "Monto (CLP)"])
    for fila in filas:
        hoja.append([fila[clave], fila["cantidad"], fila["monto"]])
    hoja.append(
        [
            "Total",
            sum(fila["cantidad"] for fila in filas),
            sum(fila["monto"] for fila in filas),
        ],
    )
    for celda in (*hoja[1], *hoja[hoja.max_row]):
        celda.font = NEGRITA
    for (celda,) in hoja.iter_rows(min_row=2, min_col=3, max_col=3):
        celda.number_format = FORMATO_CLP

    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 12
    hoja.column_dimensions["C"].width = 16
